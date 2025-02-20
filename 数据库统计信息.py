import os
import cx_Oracle
from openpyxl import load_workbook
import configparser
from datetime import datetime, timedelta

# 配置文件路径
CONFIG_FILE = 'config.ini'

def load_config():
    """从配置文件加载数据库配置信息以及SQL文件路径"""
    config = configparser.ConfigParser()
    config.read(CONFIG_FILE)

    # 加载数据库配置信息
    db_config = {
        "username": config.get('Database', 'username'),
        "password": config.get('Database', 'password'),
        "dsn": config.get('Database', 'dsn')
    }

    # 加载SQL文件路径
    sql_file_1 = config.get('SQL_Files', 'SQL_FILE_1')
    sql_file_2 = config.get('SQL_Files', 'SQL_FILE_2')

    # 加载Excel配置
    excel_file = config.get('Excel', 'EXCEL_FILE')
    sheet_name = config.get('Excel', 'SHEET_NAME')

    return db_config, sql_file_1, sql_file_2, excel_file, sheet_name

def load_sql_query(file_path, kdate, edate):
    """从SQL文件中加载查询语句并替换占位符"""
    if not os.path.exists(file_path):
        print(f"SQL文件不存在: {file_path}")
        return None
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            query = file.read().strip()
            # 替换占位符
            query = query.replace('${kdate}', kdate).replace('${edate}', edate)  # **高亮修改：替换占位符**
            return query
    except Exception as e:
        print(f"加载SQL文件失败: {e}")
        return None

def query_oracle(db_config, query, params=None):
    """执行Oracle数据库查询并返回结果"""
    try:
        connection = cx_Oracle.connect(
            db_config["username"],
            db_config["password"],
            db_config["dsn"]
        )
        cursor = connection.cursor()
        cursor.execute(query, params or {})
        result = cursor.fetchone()
        return result[0] if result else None
    except Exception as e:
        print(f"数据库查询失败: {e}")
        return None
    finally:
        if 'cursor' in locals():
            cursor.close()
        if 'connection' in locals():
            connection.close()

def update_excel(excel_file, sheet_name, cell_position, value):
    """更新Excel文件中的指定单元格"""
    try:
        workbook = load_workbook(excel_file)
        sheet = workbook[sheet_name]
        sheet[cell_position] = value
        workbook.save(excel_file)
        print(f"成功更新Excel文件: {cell_position} = {value}")
    except Exception as e:
        print(f"更新Excel文件失败: {e}")

# **新增函数：获取前一天的起始和结束时间戳**
def get_yesterday_timestamp():
    """获取前一天的起始和结束时间戳"""
    yesterday = datetime.today() - timedelta(days=1)
    # 前一天的日期时间戳：20250213000000 和 20250213235959
    data_value_start = yesterday.strftime('%Y%m%d') + '000000'  # **高亮修改：生成前一天的起始时间戳**
    data_value_end = yesterday.strftime('%Y%m%d') + '235959'    # **高亮修改：生成前一天的结束时间戳**
    return data_value_start, data_value_end, yesterday.strftime('%Y%m%d')

def save_to_new_excel(excel_file, sheet_name, yesterday_date):
    """将模版数据保存为新的Excel文件，命名为 SFM_{yesterday}.xlsx"""
    try:
        # 加载原模版文件
        workbook = load_workbook(excel_file)
        sheet = workbook[sheet_name]
        
        # 生成新的文件名
        new_file_name = f"SFM_{yesterday_date}.xlsx"
        workbook.save(new_file_name)
        print(f"新文件已保存为: {new_file_name}")
    except Exception as e:
        print(f"保存新Excel文件失败: {e}")

def main():
    """主函数"""
    db_config, sql_file_1, sql_file_2, excel_file, sheet_name = load_config()

    # **获取前一天的时间戳**
    data_value_start, data_value_end, yesterday_date = get_yesterday_timestamp()  # **高亮修改：调用函数获取前一天的时间戳**

    # 从SQL文件加载SQL查询并替换时间戳
    sql_query_1 = load_sql_query(sql_file_1, data_value_start, data_value_end)
    sql_query_2 = load_sql_query(sql_file_2, data_value_start, data_value_end)

    if not sql_query_1 or not sql_query_2:
        print("无法加载SQL查询，程序终止。")
        return

    # 执行第一条查询并更新C3单元格
    result_1 = query_oracle(db_config, sql_query_1)
    if result_1 is not None:
        update_excel(excel_file, sheet_name, "C3", result_1)
    else:
        print("未能获取到第一条查询结果，Excel未更新。")

    # 执行第二条查询并更新C4单元格
    result_2 = query_oracle(db_config, sql_query_2)
    if result_2 is not None:
        update_excel(excel_file, sheet_name, "C4", result_2)
    else:
        print("未能获取到第二条查询结果，Excel未更新。")

    # **保存为新的Excel文件**
    save_to_new_excel(excel_file, sheet_name, yesterday_date)  # **高亮修改：保存为新的文件**

if __name__ == "__main__":
    main()
